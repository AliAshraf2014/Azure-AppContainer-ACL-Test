#!/usr/bin/env python3
# Smoke-test managed-identity access to upload + download storage.

from __future__ import annotations

import argparse
import sys

from azure.core.exceptions import HttpResponseError, ServiceRequestError
from azure.identity import ManagedIdentityCredential
from azure.storage.blob import BlobServiceClient

# User-assigned managed identity (mapping agent)
MANAGED_IDENTITY_CLIENT_ID = "0ff70cda-3e59-494b-914a-146bfeceab9f"

# Upload (player) — look for a known mapping-agent output blob
UPLOAD_ACCOUNT_NAME = "s3stapplayeruks001"
UPLOAD_CONTAINER_NAME = "intellixcore-mappingagent"
UPLOAD_BLOB_PATH = (
    "TB/20260908085934/"
    "testClientA-1461933-Performance TB - 1000 rows (Tier 3 only)_trial_balance.json"
)

# Download (EDP client data) — verify read access
DOWNLOAD_ACCOUNT_NAME = "s3stedpclientdata"
DOWNLOAD_CONTAINER_NAME = "clientdata"
DOWNLOAD_ACCOUNT_URL = "https://s3stedpclientdata.blob.core.windows.net"

# Fail fast instead of hanging on private-endpoint / ACL list stalls
CONNECTION_TIMEOUT = 15
READ_TIMEOUT = 30

# Only the first slice of a blob is fetched for the content preview — enough
# for a header row, never the whole trial balance.
PREVIEW_BYTES = 64 * 1024
PREVIEW_WORDS = 40


def _log(msg: str) -> None:
    print(msg, flush=True)


def _credential() -> ManagedIdentityCredential:
    # Direct UAMI — avoids DefaultAzureCredential chain stalls.
    return ManagedIdentityCredential(client_id=MANAGED_IDENTITY_CLIENT_ID)


def _blob_service(account_name: str, account_url: str | None = None) -> BlobServiceClient:
    url = (account_url or f"https://{account_name}.blob.core.windows.net").rstrip("/")
    return BlobServiceClient(
        account_url=url,
        credential=_credential(),
        connection_timeout=CONNECTION_TIMEOUT,
        read_timeout=READ_TIMEOUT,
    )


def _words_from_bytes(data: bytes, blob_path: str, word_count: int) -> str:
    """Best-effort human-readable preview of the first bytes of a blob.

    Excel workbooks are zip archives, so decoding them as text gives noise:
    read those through openpyxl when it is available, and fall back to
    reporting the type rather than printing binary.
    """
    lower = blob_path.lower()

    if lower.endswith((".xlsx", ".xlsm")):
        try:
            import io
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            words: list[str] = []
            for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
                for cell in row:
                    if cell is None or str(cell).strip() == "":
                        continue
                    words.extend(str(cell).split())
                    if len(words) >= word_count:
                        break
                if len(words) >= word_count:
                    break
            wb.close()
            return " ".join(words[:word_count]) or "(workbook has no readable cells)"
        except ImportError:
            return "(.xlsx — install openpyxl to preview cells)"
        except Exception as exc:
            # A partial slice of a zip cannot be opened; that is expected when
            # the workbook is larger than PREVIEW_BYTES.
            return f"(.xlsx — could not parse the downloaded slice: {exc})"

    text = data.decode("utf-8-sig", errors="replace")
    if "�" in text[:200]:
        return f"(binary content — first bytes: {data[:24].hex(' ')})"
    return " ".join(text.split()[:word_count]) or "(file is empty)"


def check_upload_blob(
    account_name: str | None = None,
    container_name: str | None = None,
    blob_path: str | None = None,
) -> dict:
    """Look for the known file in the upload container."""
    account = (account_name or UPLOAD_ACCOUNT_NAME).strip()
    container = (container_name or UPLOAD_CONTAINER_NAME).strip()
    path = (blob_path or UPLOAD_BLOB_PATH).strip()

    _log("\n=== Upload container check ===")
    _log(f"Account:   {account}")
    _log(f"Container: {container}")
    _log(f"Blob:      {path}")

    client = _blob_service(account).get_blob_client(container=container, blob=path)
    try:
        _log("Getting blob properties...")
        props = client.get_blob_properties()
        row = {
            "ok": True,
            "account": account,
            "container": container,
            "blobPath": path,
            "exists": True,
            "size": props.size,
            "contentType": props.content_settings.content_type,
            "lastModified": props.last_modified.isoformat() if props.last_modified else None,
        }
        _log(f"FOUND size={props.size} contentType={props.content_settings.content_type}")
        _log(f"lastModified={row['lastModified']}")
        return row
    except Exception as exc:
        row = {
            "ok": False,
            "account": account,
            "container": container,
            "blobPath": path,
            "exists": False,
            "error": str(exc),
        }
        print(f"ERROR: {exc}", file=sys.stderr, flush=True)
        return row


def check_download_read_access(
    account_name: str | None = None,
    container_name: str | None = None,
    account_url: str | None = None,
    limit: int = 10,
) -> dict:
    """Verify the identity can list the download container.

    Uses a single lightweight list page. Full container walks hang on large /
    ACL-restricted Data Lake containers.
    """
    account = (account_name or DOWNLOAD_ACCOUNT_NAME).strip()
    container = (container_name or DOWNLOAD_CONTAINER_NAME).strip()
    url = (account_url or DOWNLOAD_ACCOUNT_URL).strip()

    _log("\n=== Download container list check ===")
    _log(f"Account:   {account}")
    _log(f"Container: {container}")
    _log(f"URL:       {url}")
    _log(f"Timeouts:  connect={CONNECTION_TIMEOUT}s read={READ_TIMEOUT}s")

    try:
        _log("Creating blob service client...")
        cc = _blob_service(account, url).get_container_client(container)

        # Lightest read probe: one page of list, capped at `limit` blobs.
        # Avoid get_container_properties + full iteration (often hangs on HNS/ACL).
        _log(f"Listing up to {limit} blob(s)...")
        pager = cc.list_blobs(results_per_page=limit).by_page()
        first_page = next(pager)
        blobs = [(b.name, getattr(b, "size", None)) for b in first_page]

        row = {
            "ok": True,
            "account": account,
            "container": container,
            "accountUrl": url,
            "canRead": True,
            "sampleBlobs": [name for name, _ in blobs],
            "sampleCount": len(blobs),
        }
        _log("LIST OK")
        for name, size in blobs:
            _log(f"  {size if size is not None else '?':>10}  {name}")
        if not blobs:
            _log("Container is listable but empty (or no blobs visible at root).")
        return row
    except StopIteration:
        # Empty container still means list permission worked.
        row = {
            "ok": True,
            "account": account,
            "container": container,
            "accountUrl": url,
            "canRead": True,
            "sampleBlobs": [],
            "sampleCount": 0,
        }
        _log("LIST OK (empty listing)")
        return row
    except (ServiceRequestError, TimeoutError, OSError) as exc:
        row = {
            "ok": False,
            "account": account,
            "container": container,
            "accountUrl": url,
            "canRead": False,
            "error": f"timeout/network: {exc}",
        }
        print(
            f"ERROR (likely firewall/private endpoint or hung list): {exc}",
            file=sys.stderr,
            flush=True,
        )
        return row
    except HttpResponseError as exc:
        row = {
            "ok": False,
            "account": account,
            "container": container,
            "accountUrl": url,
            "canRead": False,
            "error": str(exc),
            "statusCode": getattr(exc, "status_code", None),
        }
        print(f"ERROR (HTTP {row['statusCode']}): {exc}", file=sys.stderr, flush=True)
        return row
    except Exception as exc:
        row = {
            "ok": False,
            "account": account,
            "container": container,
            "accountUrl": url,
            "canRead": False,
            "error": str(exc),
        }
        print(f"ERROR: {exc}", file=sys.stderr, flush=True)
        return row


def check_download_blob(
    blob_path: str,
    account_name: str | None = None,
    container_name: str | None = None,
    account_url: str | None = None,
    word_count: int = PREVIEW_WORDS,
) -> dict:
    """Read one named file and print the first words of its content.

    Listing a container and reading a file are separate grants on an ADLS
    Gen2 account with POSIX ACLs, so a passing list check does not prove the
    identity can open a given document. This downloads real bytes.
    """
    account = (account_name or DOWNLOAD_ACCOUNT_NAME).strip()
    container = (container_name or DOWNLOAD_CONTAINER_NAME).strip()
    url = (account_url or DOWNLOAD_ACCOUNT_URL).strip()
    path = blob_path.strip()

    _log("\n=== Download blob read check ===")
    _log(f"Account:   {account}")
    _log(f"Container: {container}")
    _log(f"Blob:      {path}")

    client = _blob_service(account, url).get_blob_client(container=container, blob=path)
    try:
        _log("Getting blob properties...")
        props = client.get_blob_properties()
        _log(f"FOUND size={props.size} contentType={props.content_settings.content_type}")

        length = min(PREVIEW_BYTES, props.size or 0)
        _log(f"Downloading first {length} byte(s)...")
        data = client.download_blob(offset=0, length=length).readall() if length else b""

        preview = _words_from_bytes(data, path, word_count)
        _log(f"READ OK — first {word_count} words:")
        _log(f"  {preview}")
        return {
            "ok": True,
            "account": account,
            "container": container,
            "blobPath": path,
            "size": props.size,
            "bytesRead": len(data),
            "preview": preview,
        }
    except HttpResponseError as exc:
        print(
            f"ERROR (HTTP {getattr(exc, 'status_code', None)}): {exc}",
            file=sys.stderr,
            flush=True,
        )
        return {
            "ok": False,
            "account": account,
            "container": container,
            "blobPath": path,
            "error": str(exc),
            "statusCode": getattr(exc, "status_code", None),
        }
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr, flush=True)
        return {
            "ok": False,
            "account": account,
            "container": container,
            "blobPath": path,
            "error": str(exc),
        }


def run_smoke(
    *,
    upload: bool = True,
    download: bool = True,
    download_blob: str | None = None,
    limit: int = 10,
    words: int = PREVIEW_WORDS,
) -> list[dict]:
    """Run selected upload / download checks; return result rows."""
    results: list[dict] = []
    if upload:
        results.append(check_upload_blob())
    if download:
        results.append(check_download_read_access(limit=limit))
    if download_blob:
        results.append(check_download_blob(download_blob, word_count=words))
    return results


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Smoke-test managed-identity access to upload / download storage.",
    )
    parser.add_argument(
        "--upload",
        action="store_true",
        help=f"Check blob exists in {UPLOAD_ACCOUNT_NAME}/{UPLOAD_CONTAINER_NAME}",
    )
    parser.add_argument(
        "--download",
        action="store_true",
        help=f"List blobs in {DOWNLOAD_ACCOUNT_NAME}/{DOWNLOAD_CONTAINER_NAME}",
    )
    parser.add_argument(
        "--download-blob",
        metavar="PATH",
        help=(
            "Read this container-relative blob and print the first words of it, "
            "e.g. 'Client : 123/portal/Service : 456/documents/tb.csv'"
        ),
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=10,
        help="How many blob names to list with --download (default 10)",
    )
    parser.add_argument(
        "--words",
        type=int,
        default=PREVIEW_WORDS,
        help=f"How many words to preview with --download-blob (default {PREVIEW_WORDS})",
    )
    args = parser.parse_args(argv)

    if not args.upload and not args.download and not args.download_blob:
        parser.error("specify --upload, --download and/or --download-blob PATH")

    _log(f"Managed identity client ID: {MANAGED_IDENTITY_CLIENT_ID}")
    if args.upload:
        _log(f"Upload:   {UPLOAD_ACCOUNT_NAME}/{UPLOAD_CONTAINER_NAME}")
    if args.download or args.download_blob:
        _log(f"Download: {DOWNLOAD_ACCOUNT_NAME}/{DOWNLOAD_CONTAINER_NAME}")

    results = run_smoke(
        upload=args.upload,
        download=args.download,
        download_blob=args.download_blob,
        limit=args.limit,
        words=args.words,
    )
    errors = sum(1 for r in results if not r.get("ok"))
    _log(f"\nDone. Checks: {len(results)}, errors: {errors}")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
